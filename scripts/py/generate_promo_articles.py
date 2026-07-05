from __future__ import annotations

import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Flowable,
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "公开材料" / "推广文章"
IMG_DIR = OUT / "images"
FONT_REGULAR = Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")


ARTICLES = [
    {
        "slug": "01_customer_decision",
        "title": "别再让 AI 编程靠运气交付",
        "subtitle": "给企业客户的一份 AI 开发治理说明",
        "audience": "企业老板、研发负责人、数字化负责人",
        "hook": "AI 会写代码，但企业真正需要的是可控交付：知道规则、知道边界、知道什么时候算完成。",
        "visual_title": "从随机生成到可控交付",
        "visual_steps": ["规则入口", "任务路由", "能力单元", "验证门禁", "证据回写"],
        "sections": [
            {
                "heading": "客户真正担心的不是 AI 能不能写代码",
                "body": [
                    "现在的 AI 编码工具已经足够强，能写页面、写接口、补测试、改文档。问题在于，企业项目不是一次性生成演示代码，而是长期演进的真实系统。",
                    "在真实项目里，更大的成本来自遗忘规则、重复解释、改完不验证、不同工具各干各的，以及问题发生后找不到当时的决策依据。",
                    "Enterprise AI Development OS 的定位，就是把这些隐性成本变成可管理的工作流程。它不是替代研发团队，而是让 AI 工具按团队的交付纪律工作。",
                ],
            },
            {
                "heading": "它把项目要求变成 AI 每次都要遵守的作业单",
                "body": [
                    "规则负责告诉 AI：先看什么、不能碰什么、完成前必须验证什么。Skills 负责把高频工作沉淀成能力单元，例如规划、前端、后端、数据、测试和部署。",
                    "文档记忆负责承接跨会话任务，避免今天解释过的背景明天重新讲。审计门禁负责在公开发布、交付验收或工具切换前做结构和边界检查。",
                    "这些机制组合起来，企业得到的不是一段更漂亮的提示词，而是一套可复制、可审查、可持续改进的 AI 交付方式。",
                ],
            },
            {
                "heading": "三类业务收益",
                "bullets": [
                    "效率提升：减少重复解释、减少返工轮次，让新任务更快进入正确上下文。",
                    "风险降低：把开源边界、验证要求、工具一致性提前放进门禁，而不是发布后补救。",
                    "知识沉淀：把决策、模板、复盘和重复问题沉淀为规则、ADR、Skill 和证据链。",
                ],
            },
            {
                "heading": "适合什么企业先试",
                "body": [
                    "最适合的场景，是已经在用 AI 编程工具，但发现交付质量波动、规则容易遗忘、多人协作难以统一、老项目接手成本高的团队。",
                    "建议从 lite 模式开始：先安装入口规则、任务清单和文档模板，跑一两个真实小任务，再用公开的价值证据模板记录返工次数、验证覆盖和任务完成质量。",
                ],
            },
        ],
        "quote": "AI 写代码是起点，可控交付才是企业真正要买的结果。",
        "cta": "先用一个真实小任务试点：记录改动前后的返工次数、验证结果和规则遵守情况。",
    },
    {
        "slug": "02_investor_value",
        "title": "AI 编程的下一层机会：交付治理",
        "subtitle": "给投资人看的 Enterprise AI Development OS 价值说明",
        "audience": "投资人、创业者、技术合伙人",
        "hook": "模型能力越强，企业越需要一层可迁移的治理系统，把生成能力变成可复用的交付能力。",
        "visual_title": "业务价值飞轮",
        "visual_steps": ["规则", "Skills", "审计", "记忆", "进化"],
        "sections": [
            {
                "heading": "为什么不是又一个 AI 编程工具",
                "body": [
                    "市场上已经有很多强大的 AI 编程入口。Enterprise AI Development OS 选择站在这些工具之上，解决另一个问题：当企业同时使用多个 AI 工具时，如何让它们遵守同一套项目纪律。",
                    "它的公开版聚焦规则、Skills、文档记忆、审计门禁、安装脚本和多工具适配器。可执行规则运行时、规则命中分析、MCP 审计和团队治理面板被明确放在未来方向里，不作为当前已实现能力宣传。",
                    "这个边界很重要：它让项目先从可迁移的方法论和轻量工具切入，而不是一开始就承诺庞大的平台化能力。",
                ],
            },
            {
                "heading": "投资价值来自可复制的操作模型",
                "body": [
                    "很多 AI 研发提效停留在个人经验：某个工程师会写提示词，某个团队会调工具，某次项目里效果不错。但这类经验很难规模化。",
                    "本项目把经验拆成可复制资产：规则入口、能力单元、公开边界、审计脚本、任务模板、ADR 和价值证据模板。它让团队可以在不同项目、不同工具、不同成员之间复用同一套工作方式。",
                    "从投资视角看，这不是单点功能，而是围绕 AI 辅助开发形成的治理层。它可以服务开源社区，也可以延展到实施服务、治理支持、工具插件和团队策略包。",
                ],
            },
            {
                "heading": "商业收益应如何验证",
                "bullets": [
                    "效率：看重复解释是否减少、任务进入上下文是否更快、返工轮次是否下降。",
                    "风险：看验证门禁是否提前发现边界、结构、适配和发布问题。",
                    "知识：看一次经验能否进入规则、模板或 Skill，而不是留在聊天记录里。",
                    "成本：评估不只看首轮 token，也要把 bug 修复、返工和人工对齐成本算进去。",
                ],
            },
            {
                "heading": "当前应避免的过度承诺",
                "body": [
                    "公开材料不应宣称固定节省比例、固定缺陷降低比例，也不应暗示已经拥有完整运行时平台或团队治理面板。",
                    "更稳妥的表达是：项目已经提供可迁移的规则、Skills、审计门禁、文档记忆和适配器框架，并定义了如何用可复现实验去证明价值。",
                ],
            },
        ],
        "quote": "模型生成代码，治理系统沉淀交付能力。",
        "cta": "下一步最有价值的工作，是用脱敏小任务建立 before/after 证据，而不是先写宏大的市场百分比。",
    },
    {
        "slug": "03_open_source_community",
        "title": "给 AI 编程工具装上一套项目纪律",
        "subtitle": "面向开源社区和技术决策者的介绍",
        "audience": "开源贡献者、架构师、AI 工具使用者",
        "hook": "一个项目可以同时支持 Codex、Claude Code、Trae、Qoder、Cursor、GitHub Copilot 和 VS Code，但前提是它们读到同一套规则。",
        "visual_title": "多工具同一交付纪律",
        "visual_steps": ["Codex", "Claude", "Trae", "Qoder", "Cursor", "Copilot"],
        "sections": [
            {
                "heading": "为什么需要工具无关的规则层",
                "body": [
                    "AI 编程工具变化很快，团队不可能把项目规范绑定在某一个工具里。今天用这个插件，明天换另一个 IDE，项目纪律不能跟着丢。",
                    "Enterprise AI Development OS 把规则、Skills、文档模板、审计脚本和适配器生成器放在仓库里，让不同工具从同一个源头获得约束。",
                    "这使得 AI 辅助开发从个人习惯变成项目资产：新人进入仓库、换工具、做交接时，都能从同一套入口开始。",
                ],
            },
            {
                "heading": "公开版包含什么",
                "bullets": [
                    "规则入口：会话启动、开发顺序、目录边界、验证门禁。",
                    "Skills：规划、架构、治理、前端、后端、数据、测试和部署能力单元。",
                    "文档记忆：Backlog、总控索引、模板、ADR 和回写结构。",
                    "审计门禁：方法论结构检查、开源边界检查、就绪度评分。",
                    "工具适配：把同一套规则投放到多种 AI 编程工具。",
                ],
            },
            {
                "heading": "它不是什么",
                "body": [
                    "它不是通用 AI 应用平台，也不是完整的模型生命周期管理系统。它不承诺消除所有幻觉，也不声称当前已经实现完整的规则运行时。",
                    "它更像一层可迁移的工程治理骨架：先把项目规则、能力单元和验证方式组织好，再用审计与证据逐步增强。",
                ],
            },
            {
                "heading": "如何参与或试用",
                "body": [
                    "可以先用 lite 模式安装到一个现有项目，只加入 AI 入口规则、文档模板和任务清单骨架。完整模式会复制官方 Skills、适配器工具和审计脚本。",
                    "贡献方向可以从规则改进、Skill 补充、适配器验证、公开案例脱敏和价值证据模板开始。所有公开贡献都应通过方法论审计、开源边界检查和适配器 dry-run。",
                ],
            },
        ],
        "quote": "工具会变，项目纪律应该留在仓库里。",
        "cta": "从 lite 模式开始，把一个项目的 AI 协作入口先统一起来。",
    },
]


def ensure_dirs() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def make_visual(article: dict[str, object]) -> Path:
    path = IMG_DIR / f"{article['slug']}.png"
    w, h = 1600, 900
    img = Image.new("RGB", (w, h), "#f7f8fb")
    draw = ImageDraw.Draw(img)
    title_font = font(58, True)
    sub_font = font(30)
    label_font = font(30, True)
    small_font = font(24)

    draw.rounded_rectangle((60, 60, w - 60, h - 60), radius=28, fill="#ffffff", outline="#d8dee9", width=3)
    draw.text((110, 105), str(article["visual_title"]), font=title_font, fill="#1f2937")
    draw.text((112, 180), str(article["hook"]), font=sub_font, fill="#4b5563")

    steps = list(article["visual_steps"])
    colors_fill = ["#e8f1ff", "#eaf7ef", "#fff4df", "#f3ecff", "#e9f8fb", "#fbecef"]
    colors_line = ["#2f6fed", "#2f9e44", "#f59f00", "#7950f2", "#0c8599", "#c2255c"]
    box_w = 205 if len(steps) <= 5 else 185
    gap = 30 if len(steps) <= 5 else 22
    total = len(steps) * box_w + (len(steps) - 1) * gap
    x = (w - total) // 2
    y = 360

    for idx, step in enumerate(steps):
        left = x + idx * (box_w + gap)
        right = left + box_w
        draw.rounded_rectangle((left, y, right, y + 150), radius=22, fill=colors_fill[idx % len(colors_fill)], outline=colors_line[idx % len(colors_line)], width=4)
        label_lines = wrap_text(draw, str(step), label_font, box_w - 34)
        text_y = y + 48 - (len(label_lines) - 1) * 18
        for line in label_lines:
            tw = draw.textbbox((0, 0), line, font=label_font)[2]
            draw.text((left + (box_w - tw) / 2, text_y), line, font=label_font, fill="#111827")
            text_y += 40
        if idx < len(steps) - 1:
            ax1 = right + 6
            ax2 = right + gap - 8
            ay = y + 75
            draw.line((ax1, ay, ax2, ay), fill="#6b7280", width=4)
            draw.polygon([(ax2, ay), (ax2 - 14, ay - 9), (ax2 - 14, ay + 9)], fill="#6b7280")

    bottom = [
        ("效率", "少解释、少返工、快接手"),
        ("风险", "边界、验证、发布前置"),
        ("知识", "规则、ADR、Skills 沉淀"),
    ]
    bx = 145
    for title, desc in bottom:
        draw.rounded_rectangle((bx, 650, bx + 380, 760), radius=18, fill="#111827")
        draw.text((bx + 28, 674), title, font=label_font, fill="#ffffff")
        draw.text((bx + 28, 720), desc, font=small_font, fill="#d1d5db")
        bx += 450

    img.save(path)
    return path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Title", 24, "1F2937"),
        ("Subtitle", 13, "4B5563"),
        ("Heading 1", 16, "1F2937"),
        ("Heading 2", 13, "374151"),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def add_footer(section, text: str) -> None:
    p = section.footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(107, 114, 128)


def build_docx(article: dict[str, object], image_path: Path) -> Path:
    doc = Document()
    set_doc_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    add_footer(section, "Enterprise AI Development OS | Public promotional article | No private case data")

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(str(article["title"])).bold = True

    subtitle = doc.add_paragraph()
    subtitle.style = "Subtitle"
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(str(article["subtitle"]))

    audience = doc.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = audience.add_run(f"适读对象：{article['audience']}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_picture(str(image_path), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    hook_table = doc.add_table(rows=1, cols=1)
    hook_cell = hook_table.cell(0, 0)
    set_cell_shading(hook_cell, "EEF4FF")
    p = hook_cell.paragraphs[0]
    p.add_run(str(article["hook"])).bold = True
    p.runs[0].font.color.rgb = RGBColor(31, 41, 55)

    for section_data in article["sections"]:
        doc.add_heading(str(section_data["heading"]), level=1)
        for para in section_data.get("body", []):
            p = doc.add_paragraph(str(para))
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(6)
        for bullet in section_data.get("bullets", []):
            p = doc.add_paragraph(str(bullet), style="List Bullet")
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.space_after = Pt(4)

    quote_table = doc.add_table(rows=1, cols=1)
    cell = quote_table.cell(0, 0)
    set_cell_shading(cell, "F7F8FB")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"“{article['quote']}”")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(17, 24, 39)

    doc.add_heading("建议行动", level=1)
    p = doc.add_paragraph(str(article["cta"]))
    p.paragraph_format.line_spacing = 1.25

    doc.add_heading("公开口径边界", level=1)
    for item in [
        "本文仅使用公开仓库中的规则、Skills、审计门禁、文档记忆和价值说明材料。",
        "不包含私有策略、未脱敏案例、真实客户数据、本地路径或过程日志。",
        "不宣称固定节省比例、固定缺陷降低比例或必然交付成功。",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    path = OUT / f"{article['slug']}_{article['title']}.docx"
    doc.save(path)
    return path


class Rule(Flowable):
    def __init__(self, width: float, color=colors.HexColor("#E5E7EB")):
        super().__init__()
        self.width = width
        self.color = color

    def draw(self):
        self.canv.setStrokeColor(self.color)
        self.canv.setLineWidth(0.8)
        self.canv.line(0, 0, self.width, 0)


def register_fonts() -> tuple[str, str]:
    regular_name = "NotoSansSC"
    bold_name = "MSYHBD"
    pdfmetrics.registerFont(TTFont(regular_name, str(FONT_REGULAR)))
    pdfmetrics.registerFont(TTFont(bold_name, str(FONT_BOLD)))
    return regular_name, bold_name


def clean_pdf_text(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(article: dict[str, object], image_path: Path, regular: str, bold: str) -> Path:
    path = OUT / f"{article['slug']}_{article['title']}.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
    )
    width = A4[0] - 34 * mm
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CNTitle", fontName=bold, fontSize=24, leading=31, alignment=TA_CENTER, textColor=colors.HexColor("#111827"), spaceAfter=6))
    styles.add(ParagraphStyle("CNSub", fontName=regular, fontSize=12, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#4B5563"), spaceAfter=8))
    styles.add(ParagraphStyle("CNBody", fontName=regular, fontSize=10.5, leading=17, alignment=TA_LEFT, textColor=colors.HexColor("#1F2937"), spaceAfter=7))
    styles.add(ParagraphStyle("CNH1", fontName=bold, fontSize=15, leading=22, textColor=colors.HexColor("#111827"), spaceBefore=10, spaceAfter=7))
    styles.add(ParagraphStyle("CNQuote", fontName=bold, fontSize=13, leading=21, alignment=TA_CENTER, textColor=colors.HexColor("#111827")))
    styles.add(ParagraphStyle("CNFine", fontName=regular, fontSize=8.5, leading=13, textColor=colors.HexColor("#6B7280")))

    story = [
        Paragraph(clean_pdf_text(str(article["title"])), styles["CNTitle"]),
        Paragraph(clean_pdf_text(str(article["subtitle"])), styles["CNSub"]),
        Paragraph(clean_pdf_text(f"适读对象：{article['audience']}"), styles["CNSub"]),
        RLImage(str(image_path), width=width, height=width * 0.5625),
        Spacer(1, 8),
    ]
    hook = Table(
        [[Paragraph(clean_pdf_text(str(article["hook"])), styles["CNBody"])]],
        colWidths=[width],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF4FF")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8E6FF")),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]),
    )
    story.extend([hook, Spacer(1, 6)])

    for section_data in article["sections"]:
        story.append(Paragraph(clean_pdf_text(str(section_data["heading"])), styles["CNH1"]))
        story.append(Rule(width))
        story.append(Spacer(1, 5))
        for para in section_data.get("body", []):
            story.append(Paragraph(clean_pdf_text(str(para)), styles["CNBody"]))
        for bullet in section_data.get("bullets", []):
            story.append(Paragraph(clean_pdf_text(f"• {bullet}"), styles["CNBody"]))

    quote = Table(
        [[Paragraph(clean_pdf_text(f"“{article['quote']}”"), styles["CNQuote"])]],
        colWidths=[width],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F8FB")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ]),
    )
    story.extend([Spacer(1, 8), quote])
    story.append(Paragraph("建议行动", styles["CNH1"]))
    story.append(Paragraph(clean_pdf_text(str(article["cta"])), styles["CNBody"]))
    story.append(Paragraph("公开口径边界", styles["CNH1"]))
    for item in [
        "本文仅使用公开仓库中的规则、Skills、审计门禁、文档记忆和价值说明材料。",
        "不包含私有策略、未脱敏案例、真实客户数据、本地路径或过程日志。",
        "不宣称固定节省比例、固定缺陷降低比例或必然交付成功。",
    ]:
        story.append(Paragraph(clean_pdf_text(f"• {item}"), styles["CNFine"]))

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawCentredString(A4[0] / 2, 9 * mm, f"Enterprise AI Development OS | Public promotional article | Page {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def validate_outputs(paths: list[Path]) -> None:
    from docx import Document as ReadDocx
    from pypdf import PdfReader

    for path in paths:
        if not path.exists() or path.stat().st_size < 10_000:
            raise RuntimeError(f"Output missing or too small: {path}")
        if path.suffix.lower() == ".docx":
            doc = ReadDocx(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            if "公开口径边界" not in text or len(text) < 700:
                raise RuntimeError(f"DOCX content check failed: {path}")
        elif path.suffix.lower() == ".pdf":
            reader = PdfReader(str(path))
            if len(reader.pages) < 2:
                raise RuntimeError(f"PDF page count check failed: {path}")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if "公开口径边界" not in text or len(text) < 700:
                raise RuntimeError(f"PDF text check failed: {path}")


def main() -> None:
    ensure_dirs()
    regular, bold = register_fonts()
    outputs: list[Path] = []
    for article in ARTICLES:
        image_path = make_visual(article)
        outputs.append(image_path)
        outputs.append(build_docx(article, image_path))
        outputs.append(build_pdf(article, image_path, regular, bold))
    validate_outputs([p for p in outputs if p.suffix.lower() in {".docx", ".pdf"}])
    for path in outputs:
        rel = path.relative_to(ROOT)
        print(rel.as_posix())


if __name__ == "__main__":
    main()
